"""Kiểm thử tổng hợp nhiều chi nhánh (Mẫu 30) trên dữ liệu giả lập."""
import os
import random
import shutil
import sys
import tempfile
import threading

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from qt_app import engine  # noqa: E402
from qt_app.settings import Settings  # noqa: E402

random.seed(3)
work = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="qt_th_")
src = os.path.join(work, "data")
shutil.rmtree(src, ignore_errors=True)
os.makedirs(src)
H = ",".join("H%d" % i for i in range(21))
for br in ("5501", "5502", "5503"):
    gl = {}
    lines = [H]
    for i in range(40):
        acct = random.choice(["4211", "4221", "4232"])
        ccy = random.choice(["VND", "VND", "USD"])
        nt = random.randint(1, 900) * (100000 if ccy == "VND" else 1)
        vnd = nt if ccy == "VND" else nt * 25000
        lines.append("20251231,%s,%s,550%010d,\"Khách %d\",1%012d,S%d,%s,%s,%s,'12,20250101,20260101,4.5,20251201,0,100,DP,1,TR,'4913"
                     % (br, acct, i, i, i, i, ccy, nt, vnd))
        gl[(acct, ccy)] = gl.get((acct, ccy), 0) + nt
        gl[(acct, "000")] = gl.get((acct, "000"), 0) + vnd
    if br == "5502":
        gl[("4221", "000")] = gl.get(("4221", "000"), 0) + 5000  # chênh lệch cố ý
    with open(os.path.join(src, "%s_rt13.csv" % br), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    with open(os.path.join(src, "%s_glcb41.txt" % br), "w", encoding="utf-8") as f:
        f.write("\t".join("C%d" % i for i in range(1, 13)) + "\n")
        for (a, c), v in gl.items():
            f.write("\t".join([c, a] + ["0"] * 8 + ["0", str(v)]) + "\n")

s = Settings(path=os.path.join(work, "QT.ini"))
s.set("ngay_qt", "31/12/2025")
jobs = engine.discover_branches(src, "rt13", s)
print("Chi nhánh:", [j.brcd for j in jobs])
path, prev = engine.consolidate(s, jobs, "rt13", "5500", {"tomau": True, "kiem_tra_gl": True},
                                os.path.join(work, "Result"), log=print, cancel=threading.Event())
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(path, data_only=True)
print(wb.sheetnames)
ws = wb["Mau 30"]
for r in range(1, ws.max_row + 1):
    v = [ws.cell(row=r, column=c).value for c in range(1, 9)]
    if any(x is not None for x in v):
        print(r, v)
