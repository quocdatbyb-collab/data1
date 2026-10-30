"""Mở file kết quả bằng LibreOffice, tính lại công thức và so với giá trị đã ghi sẵn.

Dùng: python3 verify_lo.py <thư mục Result> [thư mục xuất PDF]
"""

import glob
import os
import subprocess
import sys
import time

import openpyxl
import uno
from com.sun.star.beans import PropertyValue

folder = sys.argv[1]
pdf_dir = sys.argv[2] if len(sys.argv) > 2 else None

proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore",
                         "--accept=socket,host=localhost,port=2002;urp;"],
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
ctx = None
for _ in range(60):
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        ctx = resolver.resolve("uno:socket,host=localhost,port=2002;urp;StarOffice.ComponentContext")
        break
    except Exception:  # noqa: BLE001
        time.sleep(0.5)
desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)


def prop(name, value):
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


bad = 0
checked = 0
for path in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
    wb_f = openpyxl.load_workbook(path)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(path), "_blank", 0, (prop("Hidden", True),))
    doc.calculateAll()
    for ws in wb_f.worksheets:
        sheet = doc.Sheets.getByName(ws.title)
        wv = wb_v[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    lo = sheet.getCellByPosition(cell.column - 1, cell.row - 1)
                    got = lo.getValue()
                    exp = wv.cell(row=cell.row, column=cell.column).value
                    if exp is None:
                        continue  # mẫu tuỳ chỉnh (openpyxl) không ghi sẵn giá trị – Excel tự tính khi mở
                    checked += 1
                    try:
                        exp = float(exp)
                    except (TypeError, ValueError):
                        exp = 0
                    if abs(got - exp) > 0.01 or lo.getError() != 0:
                        bad += 1
                        if bad < 30:
                            print("SAI %s!%s%d %s: LO=%s, ghi sẵn=%s" % (
                                os.path.basename(path), ws.title, cell.row, cell.value, got, exp))
    if pdf_dir:
        os.makedirs(pdf_dir, exist_ok=True)
        out = os.path.join(pdf_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        doc.storeToURL(uno.systemPathToFileUrl(out), (prop("FilterName", "calc_pdf_Export"),))
    doc.close(True)
print("Đã kiểm tra %d công thức, sai %d" % (checked, bad))
try:
    desktop.terminate()
except Exception:  # noqa: BLE001
    pass
proc.wait(timeout=30)
